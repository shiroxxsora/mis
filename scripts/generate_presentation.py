#!/usr/bin/env python3
"""Generate MIS project presentation (PPTX) with app color palette and Mermaid diagrams."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PRES_DIR = ROOT / "docs" / "presentation"
DIAGRAMS = PRES_DIR / "diagrams"
BUILD = PRES_DIR / "build"
OUTPUT = PRES_DIR / "MIS-presentation.pptx"
MERMAID_CONFIG = PRES_DIR / "mermaid-config.json"

# Palette from frontend/src/index.css
C_PRIMARY = RGBColor(0xFF, 0x4B, 0x23)
C_PRIMARY_HOVER = RGBColor(0xE2, 0x42, 0x1F)
C_PRIMARY_LIGHT = RGBColor(0xFF, 0xE1, 0xD7)
C_SIDEBAR = RGBColor(0x2A, 0x2A, 0x2E)
C_SIDEBAR_TEXT = RGBColor(0xE8, 0xE8, 0xEA)
C_SIDEBAR_MUTED = RGBColor(0xA0, 0xA0, 0xA8)
C_TABLE_HEADER = RGBColor(0xEE, 0xF0, 0xF3)
C_TABLE_STRIPE = RGBColor(0xF8, 0xF8, 0xF9)
C_TEXT = RGBColor(0x11, 0x11, 0x11)
C_MUTED = RGBColor(0x66, 0x66, 0x66)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_BORDER = RGBColor(0xD8, 0xD8, 0xD8)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def render_mermaid() -> dict[str, Path]:
    BUILD.mkdir(parents=True, exist_ok=True)
    images: dict[str, Path] = {}
    for mmd in sorted(DIAGRAMS.glob("*.mmd")):
        out = BUILD / f"{mmd.stem}.png"
        cmd = [
            "npx",
            "--yes",
            "@mermaid-js/mermaid-cli@11.4.0",
            "-i",
            str(mmd),
            "-o",
            str(out),
            "-c",
            str(MERMAID_CONFIG),
            "-b",
            "white",
            "-w",
            "2400",
            "-H",
            "1350",
            "-s",
            "2",
        ]
        print(f"Rendering {mmd.name} …")
        # Windows: npx.cmd не находится без shell=True
        use_shell = sys.platform == "win32"
        subprocess.run(cmd, cwd=ROOT, check=True, shell=use_shell)
        images[mmd.stem] = out
    return images


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title: str) -> None:
    bar = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(0),
        Inches(0),
        SLIDE_W,
        Inches(0.95),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_SIDEBAR
    bar.line.fill.background()
    tf = bar.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = "Segoe UI"
    tf.margin_left = Inches(0.45)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    accent = slide.shapes.add_shape(1, Inches(0), Inches(0.95), SLIDE_W, Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = C_PRIMARY
    accent.line.fill.background()


def add_bullets(slide, items: list[str], top: float = 1.25, left: float = 0.55, width: float = 12.2) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.font.name = "Segoe UI"
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(8)


def _image_fit_inches(image_path: Path, max_w_in: float, max_h_in: float) -> tuple[float, float]:
    """Return (width_in, height_in) preserving aspect ratio inside max_w × max_h."""
    with Image.open(image_path) as im:
        w_px, h_px = im.size
    if w_px <= 0 or h_px <= 0:
        return max_w_in, max_h_in
    aspect = w_px / h_px
    # width-limited candidate
    w1, h1 = max_w_in, max_w_in / aspect
    if h1 <= max_h_in:
        return w1, h1
    return max_h_in * aspect, max_h_in


def add_table_slide(
    prs: Presentation,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float] | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, title)

    n_rows = len(rows) + 1
    n_cols = len(headers)
    left, top, width, height = Inches(0.45), Inches(1.2), Inches(12.4), Inches(5.9)
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table

    if col_widths:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Inches(w)

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_TABLE_HEADER
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(13)
            p.font.name = "Segoe UI"
            p.font.color.rgb = C_TEXT

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            if r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_TABLE_STRIPE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.name = "Segoe UI"
                p.font.color.rgb = C_TEXT


def add_image_slide(
    prs: Presentation,
    title: str,
    image: Path,
    caption: str | None = None,
    *,
    intro: str | None = None,
    max_pic_w: float = 12.4,
    max_pic_h: float | None = None,
) -> None:
    """Insert diagram without stretching (aspect ratio preserved)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, title)

    top_after_header = 1.02
    if intro:
        intro_box = slide.shapes.add_textbox(Inches(0.45), Inches(top_after_header), Inches(12.4), Inches(1.35))
        itf = intro_box.text_frame
        itf.word_wrap = True
        for i, line in enumerate(intro.split("\n")):
            p = itf.paragraphs[0] if i == 0 else itf.add_paragraph()
            p.text = line
            p.font.size = Pt(13)
            p.font.name = "Segoe UI"
            p.font.color.rgb = C_TEXT
            p.space_after = Pt(4)
        top_after_header += 1.42

    # Reserve space for optional caption at bottom
    caption_h = 0.42 if caption else 0.12
    default_max_h = 7.45 - top_after_header - caption_h - 0.08
    max_h = max_pic_h if max_pic_h is not None else default_max_h

    w_in, h_in = _image_fit_inches(image, max_pic_w, max_h)
    left_in = (13.333 - w_in) / 2.0
    pic_top = top_after_header + 0.05
    slide.shapes.add_picture(str(image), Inches(left_in), Inches(pic_top), width=Inches(w_in), height=Inches(h_in))

    if caption:
        cap_top = min(7.05, pic_top + h_in + 0.06)
        cap = slide.shapes.add_textbox(Inches(0.45), Inches(cap_top), Inches(12.4), Inches(0.4))
        tf = cap.text_frame
        tf.text = caption
        p = tf.paragraphs[0]
        p.font.size = Pt(11)
        p.font.color.rgb = C_MUTED
        p.font.name = "Segoe UI"
        p.alignment = PP_ALIGN.CENTER


def build_presentation(images: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 — Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_SIDEBAR)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(6.85), SLIDE_W, Inches(0.65))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_PRIMARY
    bar.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.2))
    tf = title.text_frame
    tf.text = "MIS"
    p = tf.paragraphs[0]
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = "Segoe UI"

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.15), Inches(11.5), Inches(1.5))
    stf = sub.text_frame
    stf.text = "Медицинская информационная система\nСтоматологическая клиника · Docker · GraphQL · AI"
    for para in stf.paragraphs:
        para.font.size = Pt(22)
        para.font.color.rgb = C_SIDEBAR_MUTED
        para.font.name = "Segoe UI"

    # 2 — About
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "О системе")
    add_bullets(
        slide,
        [
            "Единый стек для реестра пациентов, записи на приём и ведения визита",
            "Веб-интерфейс React с входом через Keycloak (OIDC, PKCE)",
            "Данные клиники в PostgreSQL; доступ через Hasura GraphQL с ролевой моделью",
            "Оркестрация LLM, распознавания снимков и файлов — service-logic",
            "Одна клиника: общий реестр для всех авторизованных сотрудников",
        ],
    )

    # 3 — Stack
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Технологический стек")
    stacks = [
        ("Frontend", "React, Vite, TypeScript, nginx"),
        ("API", "Spring Cloud Gateway, Spring Boot"),
        ("Данные", "Hasura, PostgreSQL 18, Liquibase"),
        ("Auth", "Keycloak 24, JWT RS256"),
        ("ML / AI", "FastAPI, DentalNet, внешний LLM"),
        ("Файлы", "MinIO (S3), бакет mis-patient-images"),
        ("Инфра", "Docker Compose"),
    ]
    x0, y0 = 0.45, 1.25
    for i, (label, desc) in enumerate(stacks):
        col = i % 2
        row = i // 2
        x = x0 + col * 6.35
        y = y0 + row * 1.45
        card = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(6.0), Inches(1.25))
        card.fill.solid()
        card.fill.fore_color.rgb = C_PRIMARY_LIGHT if i % 2 == 0 else C_TABLE_STRIPE
        card.line.color.rgb = C_PRIMARY
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.text = f"{label}\n{desc}"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = C_TEXT
        tf.paragraphs[0].font.name = "Segoe UI"
        if len(tf.paragraphs) > 1:
            tf.paragraphs[1].font.size = Pt(13)
            tf.paragraphs[1].font.color.rgb = C_MUTED
            tf.paragraphs[1].font.name = "Segoe UI"

    # 4 — User journey
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Как пользователь работает в системе")
    add_bullets(
        slide,
        [
            "1) Вход через Keycloak (OIDC) и получение роли: врач или регистратура",
            "2) Работа с пациентами и приёмами через live-данные GraphQL (Hasura subscriptions)",
            "3) Ведение приёма через REST: статусы, заметки, отмена, сохранение визита",
            "4) AI-функции для врача: распознавание снимка (DentalNet), саммари и прогноз явки (LLM)",
            "5) Документы: направление (врач) или чек (регистратура), печать из UI",
        ],
        top=1.3,
    )

    # 5 — Core architecture
    add_image_slide(
        prs,
        "Архитектура (Docker Compose)",
        images["architecture"],
        "1 слайд вместо нескольких: полный контур сервисов, портов и связей.",
    )

    # 6 — Appointment lifecycle
    add_image_slide(
        prs,
        "Жизненный цикл приёма",
        images["fsm"],
        caption=(
            "При одновременном сохранении: optimistic lock в БД → HTTP 409 Conflict, если статус уже изменился."
        ),
        intro=(
            "Автомат состояний (англ. FSM, finite state machine) — набор допустимых статусов приёма "
            "и правил перехода между ними. Нельзя «перепрыгнуть» из запланированного сразу в завершённый: "
            "смена статуса только по стрелкам на схеме. В Hasura прямой update записей приёмов закрыт; "
            "обработка визита идёт через REST API с проверкой этих правил на сервере."
        ),
    )

    # 7 — DB versioning
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Версионирование БД (Liquibase)")
    add_bullets(
        slide,
        [
            "Схема БД меняется только через changelog в backend/liquibase",
            "Каждая миграция — отдельный SQL-файл, порядок фиксирован и воспроизводим",
            "При старте app-db контейнер применяет миграции автоматически",
            "Это убирает «дрейф схемы» между разработкой, тестом и демо",
            "Hasura metadata опирается на актуальную схему после миграций",
        ],
        top=1.3,
    )

    # 8 — Testing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Тестирование и качество")
    add_bullets(
        slide,
        [
            "Backend Java: unit/web-тесты для gateway и service-logic",
            "Покрыты ключевые кейсы: FSM переходов, контроллеры, обработка ошибок, парсинг LLM",
            "Recognition-service: Python-тесты inference и API-эндпоинтов",
            "Тесты запускаются локально и в CI перед релизом",
            "Подход: валидация бизнес-правил + защита от регрессий",
        ],
        top=1.3,
    )

    # 9 — Roles and docs flow
    add_image_slide(
        prs,
        "Роли и сценарии (врач / регистратура)",
        images["use-cases"],
        "Документы: referral для врача, receipt для регистратуры.",
    )

    # 10 — AI
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Искусственный интеллект")
    add_bullets(
        slide,
        [
            "DentalNet (recognition-service): классификация снимков, опционально SHAP/LIME",
            "LLM (вне Docker): саммари истории визитов, прогноз вероятности явки 0–100%",
            "Backfill при старте: пересчёт attendance_probability для пациентов без значения",
            "Все вызовы AI — через service-logic; gateway проверяет JWT пользователя",
        ],
        top=1.3,
    )

    # 11 — Security
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Безопасность")
    add_bullets(
        slide,
        [
            "JWT на api-gateway (issuer, azp=mis-frontend); service-logic только в Docker-сети",
            "Hasura: роли user / registrar, x-hasura-role из JWT; update appointments закрыт",
            "file_path снимков скрыт в GraphQL; скачивание — REST с JWT",
            "Секреты в .env: HASURA_ADMIN_SECRET, пароли Keycloak/MinIO",
            "Production: сменить дефолты, HASURA_GRAPHQL_DEV_MODE=false",
        ],
        top=1.3,
    )

    # 12 — Run
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_WHITE)
    add_header_bar(slide, "Запуск и демо")
    add_bullets(
        slide,
        [
            "cp .env.example .env && docker compose up -d --build",
            "sh scripts/keycloak-create-demo-user.sh",
            "sh scripts/keycloak-create-registrar-user.sh",
            "UI: http://localhost:3000 · API: http://localhost:8084",
            "Документация: docs/architecture.md, docs/use-cases.md",
        ],
        top=1.3,
    )

    # End
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_SIDEBAR)
    end = slide.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11), Inches(2))
    etf = end.text_frame
    etf.text = "Спасибо за внимание"
    etf.paragraphs[0].font.size = Pt(40)
    etf.paragraphs[0].font.bold = True
    etf.paragraphs[0].font.color.rgb = C_WHITE
    etf.paragraphs[0].font.name = "Segoe UI"
    etf.paragraphs[0].alignment = PP_ALIGN.CENTER

    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


def main() -> int:
    if not MERMAID_CONFIG.exists():
        print("Missing mermaid config", file=sys.stderr)
        return 1
    images = render_mermaid()
    build_presentation(images)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
